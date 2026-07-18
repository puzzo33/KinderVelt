"""KV-FOS test suite.

Covers the requirements a reviewer would check first: idempotent
reprocessing without duplicates, supersession on re-upload, gap detection
and recovery, reconciliation breaks, vendor learning, auto-explained
resolutions, materiality flags, the lock/amend lifecycle, and the
reminder state machine.
"""

import json
import os
import shutil
from datetime import date
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from kvfos import cli, lock as lock_mod, pipeline
from kvfos.config import load_knowledge
from make_sample_data import build
from tests.conftest import MONTH


def close(root, month=MONTH, **kw):
    return pipeline.close_month(root, month, **kw)


def read_json(root, month, name):
    with open(root / "months" / month / "derived" / name, encoding="utf-8") as f:
        return json.load(f)


def findings_by_rule(root, month=MONTH):
    out = {}
    for f in read_json(root, month, "exceptions.json"):
        out.setdefault(f["rule"], []).append(f)
    return out


# ---------------------------------------------------------------------------
# happy path

def test_full_close_happy_path(root):
    s = close(root)
    assert s["fully_reconciled"]
    assert s["gaps_blocking"] == 0
    assert s["transactions"] == 9
    assert s["confidence"] == 100
    traces = read_json(root, MONTH, "traces.json")
    assert len(traces) == 1 and traces[0]["chain_complete"]
    # realized Wise rate takes precedence (§9.9): ledger income is Verified USD
    txns = [json.loads(l) for l in
            open(root / "months" / MONTH / "derived" / "transactions.jsonl",
                 encoding="utf-8")]
    income = [t for t in txns if t["direction"] == "income"
              and t["amount_uah"] == "498000.00"][0]
    assert income["fx_source"].startswith("wise-realized")
    assert income["usd_status"] == "verified"
    assert income["funding_source"] == "fokv_general"


def test_expected_exceptions(root):
    close(root)
    rules = findings_by_rule(root)
    assert "new_vendor" in rules            # ТОВ Арт-Плюс not in vendor master
    assert "large_one_time_expense" in rules
    # everything bank-matched, so no unmatched findings
    assert "unmatched_bank_transaction" not in rules
    assert "unmatched_transfer" not in rules


# ---------------------------------------------------------------------------
# idempotency & dedupe (the core re-upload requirement)

def test_reprocess_is_idempotent(root):
    close(root)
    txns1 = (root / "months" / MONTH / "derived" / "transactions.jsonl").read_text()
    reg1 = (root / "registry" / "master_register.csv").read_text()
    close(root)
    txns2 = (root / "months" / MONTH / "derived" / "transactions.jsonl").read_text()
    reg2 = (root / "registry" / "master_register.csv").read_text()
    assert txns1 == txns2
    assert reg1 == reg2
    assert reg1.count("\n") == 10  # header + 9 txns, no dupes


def test_reupload_supersedes_not_duplicates(root):
    inputs = root / "months" / MONTH / "inputs"
    original = inputs / f"bank_privat_uah_{MONTH}.csv"
    revised = inputs / f"bank_privat_uah_{MONTH}_v2.csv"
    shutil.copy(original, revised)
    # make the revised file clearly newer
    os.utime(original, (1, 1))
    s = close(root)
    docs = read_json(root, MONTH, "documents.json")
    bank_docs = [d for d in docs if d["doc_type"] == "bank_statement"
                 and (d["account"] == "privat_uah")]
    assert len(bank_docs) == 2
    superseded = [d for d in bank_docs if d["superseded_by"]]
    assert len(superseded) == 1
    assert superseded[0]["file"] == original.name
    # bank lines counted once, not twice
    assert s["bank_lines"] == 9
    assert s["fully_reconciled"]


# ---------------------------------------------------------------------------
# gaps and the re-upload loop

def test_missing_statement_gap_then_recovery(root):
    inputs = root / "months" / MONTH / "inputs"
    stmt = inputs / f"bank_oschad_uah_{MONTH}.csv"
    saved = stmt.read_bytes()
    stmt.unlink()

    s = close(root)
    assert s["gaps_blocking"] >= 1
    gaps = read_json(root, MONTH, "gaps.json")
    assert any(g["kind"] == "missing_input" and "oschad" in g["title"].lower()
               for g in gaps)
    assert s["confidence"] < 100
    assert "missing_required_input" in findings_by_rule(root)
    # cannot approve a degraded month
    with pytest.raises(lock_mod.LockError):
        lock_mod.approve(root / "months" / MONTH, MONTH)

    stmt.write_bytes(saved)   # treasurer re-uploads
    s = close(root)
    assert s["gaps_blocking"] == 0 and s["fully_reconciled"]
    assert s["confidence"] == 100


def test_us_side_statement_and_trace_link(root):
    close(root)
    traces = read_json(root, MONTH, "traces.json")
    us_link = [l for l in traces[0]["links"] if l["step"] == "us_bank_debit"][0]
    assert us_link["status"] == "matched"          # $12,000 left TD Bank
    assert traces[0]["chain_complete"]
    # US statement absent -> material gap (not blocking), chain incomplete
    (root / "months" / MONTH / "inputs" / f"bank_td_usd_{MONTH}.csv").unlink()
    s = close(root)
    assert s["gaps_blocking"] == 0
    gaps = read_json(root, MONTH, "gaps.json")
    assert any("US bank statement" in g["title"] for g in gaps)
    traces = read_json(root, MONTH, "traces.json")
    us_link = [l for l in traces[0]["links"] if l["step"] == "us_bank_debit"][0]
    assert us_link["status"] == "missing"
    assert not traces[0]["chain_complete"]
    assert traces[0]["deposit_confirmed"]          # UA side still verified


def test_corrupt_document_becomes_gap_not_crash(root):
    wb_path = root / "months" / MONTH / "inputs" / f"workbook_{MONTH}.xlsx"
    wb_path.write_bytes(b"this is not an excel file")
    s = close(root)   # must not raise
    assert s["gaps_blocking"] >= 1
    gaps = read_json(root, MONTH, "gaps.json")
    assert any(g["kind"] == "unreadable_document" for g in gaps)


def test_missing_fx_rates_flagged_not_fabricated(root):
    (root / "knowledge" / "fx_rates.yaml").write_text(
        "source: NBU\nfallback_policy: nearest_prior_within_7_days\nrates: {}\n",
        encoding="utf-8")
    close(root)
    gaps = read_json(root, MONTH, "gaps.json")
    assert any(g["kind"] == "missing_fx_rate" for g in gaps)
    txns = [json.loads(l) for l in
            open(root / "months" / MONTH / "derived" / "transactions.jsonl",
                 encoding="utf-8")]
    # ledger income keeps the realized Wise rate; everything else Missing USD
    assert all(t["usd_status"] == "missing" for t in txns
               if not (t.get("fx_source") or "").startswith("wise-realized"))


# ---------------------------------------------------------------------------
# reconciliation

def test_reconciliation_break_detected(root):
    wb_path = root / "months" / MONTH / "inputs" / f"workbook_{MONTH}.xlsx"
    wb = load_workbook(wb_path)
    ws = wb["Balances"]
    ws.cell(row=2, column=3, value=999999)  # tamper privat_uah closing
    wb.save(wb_path)
    s = close(root)
    assert not s["fully_reconciled"]
    rules = findings_by_rule(root)
    assert "reconciliation_break" in rules
    assert s["confidence"] < 100


def test_duplicate_payment_detected(root):
    wb_path = root / "months" / MONTH / "inputs" / f"workbook_{MONTH}.xlsx"
    wb = load_workbook(wb_path)
    ws = wb["GeneralLedger"]
    ws.append([f"{MONTH}-18", "privat_uah", "631-SHIP", "Нова Пошта",
               "Доставка матеріалів до Харкова", -3200, "Kharkiv", "EXP-77"])
    wb.save(wb_path)
    close(root)
    dups = findings_by_rule(root)["duplicate_payment"]
    assert len(dups) == 1 and dups[0]["severity"] == "high"
    assert len(dups[0]["txn_ids"]) == 2


# ---------------------------------------------------------------------------
# learning (SPEC §7.6)

def test_learned_alias_resolves_new_vendor(root):
    alias_file = root / "knowledge" / "learned" / "vendor_aliases.yaml"
    alias_file.write_text(yaml.safe_dump({
        "aliases": [{"alias": "ТОВ Арт-Плюс", "vendor": "nova_poshta",
                     "confirmed": "2026-07-01", "by": "treasurer"}]}),
        encoding="utf-8")
    close(root)
    rules = findings_by_rule(root)
    assert "new_vendor" not in rules


def test_learned_resolution_auto_explains(root):
    res_file = root / "knowledge" / "learned" / "resolutions.yaml"
    res_file.write_text(yaml.safe_dump({
        "resolutions": [{"rule": "large_one_time_expense",
                         "vendor": "kv_payroll",
                         "explanation": "Monthly payroll run — expected size.",
                         "accepted": "2026-07-01", "by": "treasurer"}]}),
        encoding="utf-8")
    s = close(root)
    assert s["auto_explained"] >= 1
    payroll = [f for f in read_json(root, MONTH, "exceptions.json")
               if f["rule"] == "large_one_time_expense"
               and "Заробітна" in f["title"]][0]
    assert payroll["status"] == "auto_explained"
    assert "payroll run" in payroll["resolution"]


# ---------------------------------------------------------------------------
# lock / amend lifecycle (SPEC §9.8)

def test_approve_lock_and_amend(root):
    close(root)
    lock_mod.approve(root / "months" / MONTH, MONTH)
    with pytest.raises(pipeline.CloseError):
        close(root)  # locked without --amend
    # accountant restates: extra expense appears
    wb_path = root / "months" / MONTH / "inputs" / f"workbook_{MONTH}.xlsx"
    wb = load_workbook(wb_path)
    wb["GeneralLedger"].append(
        [f"{MONTH}-28", "privat_uah", "685-ACCT", "ТОВ Аудит-Сервіс",
         "Бухгалтерські послуги", -5000, "", "INV-200"])
    wb["Balances"].cell(row=2, column=3, value=431750)
    wb.save(wb_path)
    s = close(root, amend=True)
    assert s["amended"]
    amendment = (root / "months" / MONTH / "reports" / "AMENDMENT.md").read_text(
        encoding="utf-8")
    assert "₴-5,000.00" in amendment or "₴5,000.00" in amendment


# ---------------------------------------------------------------------------
# reminder / status state machine

def test_month_status_lifecycle(root, monkeypatch):
    monkeypatch.setattr(cli, "date",
                        type("D", (), {"today": staticmethod(lambda: date(2026, 7, 15)),
                                       "fromisoformat": date.fromisoformat}))
    assert cli.months_since_start(root) == ["2026-06"]
    assert cli.month_status(root, MONTH) == "inputs_ready"
    close(root)
    assert cli.month_status(root, MONTH) == "closed"
    # treasurer replaces a file -> stale
    stats = root / "months" / MONTH / "inputs" / f"CenterUpdates_{MONTH}.csv"
    stats.write_text(stats.read_text(encoding="utf-8") +
                     "Kropyvnytskyi,0,0,0,0\n", encoding="utf-8")
    assert cli.month_status(root, MONTH) == "stale_inputs_changed"
    close(root)
    lock_mod.approve(root / "months" / MONTH, MONTH)
    assert cli.month_status(root, MONTH) == "approved"
    # a month with no folder at all
    assert cli.month_status(root, "2026-05") == "awaiting_inputs"


def test_remind_github_issue_format(root, capsys):
    rc = cli.main(["--root", str(root), "remind", "--format", "github-issue"])
    assert rc == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["actionable"] >= 0
    assert "monthly" in payload["title"].lower()


# ---------------------------------------------------------------------------
# knowledge base validation

def test_invalid_knowledge_rejected(root):
    bad = root / "knowledge" / "vendors.yaml"
    data = yaml.safe_load(bad.read_text(encoding="utf-8"))
    data["vendors"][0]["category"] = "not_a_category"
    bad.write_text(yaml.safe_dump(data, allow_unicode=True), encoding="utf-8")
    from kvfos.config import ConfigError
    with pytest.raises(ConfigError):
        load_knowledge(root)
