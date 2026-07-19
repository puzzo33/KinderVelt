"""Document registry (SPEC §9 Stage 1).

Content-addressed and idempotent:

* A document's identity is the SHA-256 of its bytes. Re-importing the same
  file yields the same doc_id — no duplicates, ever.
* Document types that must be unique per key (one accounting workbook per
  month, one statement per bank account per month, one stats report per
  month) are deduplicated by *supersession*: when several candidates exist,
  the newest file wins and the others are marked superseded (retained in
  the registry history, excluded from processing). Replacing a file in the
  inputs folder and re-running `close` therefore reprocesses cleanly.
* The registry snapshot is persisted per month and merged into a cumulative
  registry across all months.
"""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

import yaml

from .model import DocType, Document

# doc types where only one live document may exist per uniqueness key
UNIQUE_KEYS = {
    DocType.ACCOUNTING_WORKBOOK.value: lambda d: "workbook",
    DocType.CENTER_STATS.value: lambda d: "stats",
    DocType.BANK_STATEMENT.value: lambda d: f"bank:{d.account or 'unknown'}",
    DocType.PLATFORM_EXPORT.value: lambda d: f"platform:{d.account or 'unknown'}",
}

_PLATFORM_RE = re.compile(r"(zeffy|stripe|paypal)", re.I)

_FILENAME_HINTS = [
    (re.compile(r"wise", re.I), DocType.WISE_TRANSFER),
    (_PLATFORM_RE, DocType.PLATFORM_EXPORT),
    (re.compile(r"(bank|statement|виписк)", re.I), DocType.BANK_STATEMENT),
    (re.compile(r"(centerupdate|stats|statistic|metrics)", re.I), DocType.CENTER_STATS),
    (re.compile(r"(funding.?request|request.?letter)", re.I), DocType.FUNDING_REQUEST),
    (re.compile(r"(invoice|рахунок-фактура)", re.I), DocType.INVOICE),
    (re.compile(r"(workbook|accounting|ledger|облік)", re.I), DocType.ACCOUNTING_WORKBOOK),
]

_BANK_ACCOUNT_RE = re.compile(r"bank[_-]([a-z0-9_]+?)[_-]\d{4}-\d{2}", re.I)


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def detect_type(path: Path, manifest: dict) -> tuple[str, str | None]:
    """Return (doc_type, account). Manifest entries override heuristics."""
    entry = manifest.get(path.name, {})
    if "type" in entry:
        return entry["type"], entry.get("account")

    name = path.name
    account = None
    m = _BANK_ACCOUNT_RE.search(name)
    if m:
        account = m.group(1)
    for rx, dtype in _FILENAME_HINTS:
        if rx.search(name):
            if dtype is DocType.PLATFORM_EXPORT:
                account = _PLATFORM_RE.search(name).group(1).lower()
            return dtype.value, account
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        # an Excel file with ledger-like sheets is the workbook
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True)
            sheets = " ".join(s.lower() for s in wb.sheetnames)
            wb.close()
            if any(w in sheets for w in ("ledger", "журнал", "payroll", "зарпл")):
                return DocType.ACCOUNTING_WORKBOOK.value, None
        except Exception:
            pass
    return DocType.UNKNOWN.value, account


def scan_inputs(month_dir: Path, month: str) -> list[Document]:
    """Scan a month's inputs folder into Document records with
    supersession applied."""
    inputs = month_dir / "inputs"
    manifest = {}
    manifest_path = inputs / "manifest.yaml"
    if manifest_path.exists():
        with open(manifest_path, encoding="utf-8") as f:
            manifest = (yaml.safe_load(f) or {}).get("files", {})

    docs: list[Document] = []
    if not inputs.exists():
        return docs

    # portal uploads of binary files travel as base64 text (<name>.b64);
    # materialize the real file before scanning so it behaves normally
    import base64
    for b64_path in inputs.rglob("*.b64"):
        target = b64_path.with_suffix("")
        try:
            decoded = base64.b64decode(b64_path.read_text(), validate=True)
        except Exception:
            continue
        if not target.exists() or target.read_bytes() != decoded:
            target.write_bytes(decoded)

    for path in sorted(inputs.rglob("*")):
        if (not path.is_file()
                or path.name.startswith(".")       # portal markers, .gitkeep
                or path.suffix == ".b64"
                or path.name in ("manifest.yaml", "CHECKLIST.md")):
            continue
        dtype, account = detect_type(path, manifest)
        sha = sha256_file(path)
        docs.append(Document(
            doc_id=sha[:12],
            file=str(path.relative_to(inputs)),
            doc_type=dtype,
            month=month,
            account=account,
            sha256=sha,
            size=path.stat().st_size,
        ))

    _apply_supersession(docs, inputs)
    return docs


def _apply_supersession(docs: list[Document], inputs: Path) -> None:
    """Among live documents sharing a uniqueness key, the newest file
    (mtime, then name) wins; the rest are superseded."""
    groups: dict[str, list[Document]] = {}
    for d in docs:
        keyfn = UNIQUE_KEYS.get(d.doc_type)
        if keyfn:
            groups.setdefault(keyfn(d), []).append(d)
    for key, group in groups.items():
        if len(group) <= 1:
            continue
        group.sort(key=lambda d: ((inputs / d.file).stat().st_mtime, d.file))
        winner = group[-1]
        for d in group[:-1]:
            d.superseded_by = winner.doc_id
            d.notes = f"superseded by newer {d.doc_type} ({winner.file})"


def live(docs: list[Document], doc_type: str | None = None) -> list[Document]:
    out = [d for d in docs if d.superseded_by is None]
    if doc_type:
        out = [d for d in out if d.doc_type == doc_type]
    return out


def save_registry(docs: list[Document], month_dir: Path, root: Path) -> None:
    """Persist the month snapshot and merge into the cumulative registry."""
    derived = month_dir / "derived"
    derived.mkdir(parents=True, exist_ok=True)
    with open(derived / "documents.json", "w", encoding="utf-8") as f:
        json.dump([d.as_dict() for d in docs], f, ensure_ascii=False, indent=2)

    reg_dir = root / "registry"
    reg_dir.mkdir(parents=True, exist_ok=True)
    cum_path = reg_dir / "documents.json"
    cumulative: dict[str, dict] = {}
    if cum_path.exists():
        with open(cum_path, encoding="utf-8") as f:
            cumulative = {d["doc_id"]: d for d in json.load(f)}
    # replace this month's entries wholesale (idempotent reprocess)
    cumulative = {k: v for k, v in cumulative.items()
                  if v["month"] != docs[0].month} if docs else cumulative
    for d in docs:
        cumulative[d.doc_id] = d.as_dict()
    with open(cum_path, "w", encoding="utf-8") as f:
        json.dump(sorted(cumulative.values(), key=lambda d: (d["month"], d["file"])),
                  f, ensure_ascii=False, indent=2)
