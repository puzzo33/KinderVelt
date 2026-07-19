"""Stage 1 — document readers (SPEC §9 Stage 1, §6.3).

Readers accept Excel, CSV, YAML and PDF (native text). Scanned PDFs with
no text layer are not silently guessed at: the document is marked
low-confidence and a gap is raised asking for OCR or a data re-upload.
Every extracted row keeps a locator so figures remain traceable (§10.6).
"""

from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from pathlib import Path

import yaml

from .model import Document, SourceRef, money

# ---------------------------------------------------------------------------
# header synonyms (EN / UA), lowercase
_HEADER_SYNONYMS = {
    "date": {"date", "дата"},
    "bank_account": {"bankaccount", "account", "рахунок банку", "банк"},
    "account_code": {"accountcode", "code", "стаття", "рахунок обліку", "коа"},
    "counterparty": {"counterparty", "vendor", "контрагент", "постачальник"},
    "description": {"description", "purpose", "призначення", "опис"},
    "amount_uah": {"amountuah", "amount", "сума", "сума грн"},
    "center": {"center", "centre", "центр"},
    "invoice_ref": {"invoiceref", "invoice", "документ", "рахунок-фактура"},
    "balance": {"balance", "залишок"},
    "employee": {"employee", "піб", "працівник"},
    "role": {"role", "посада"},
    "gross_uah": {"grossuah", "нараховано"},
    "taxes_uah": {"taxesuah", "taxes", "податки"},
    "net_uah": {"netuah", "до виплати"},
    "opening_uah": {"openinguah", "opening", "залишок на початок"},
    "closing_uah": {"closinguah", "closing", "залишок на кінець"},
    "children_served": {"childrenserved", "children", "діти"},
    "consultations": {"consultations", "консультації"},
    "classes": {"classes", "заняття"},
    "new_children": {"newchildren", "нові діти"},
    "gross": {"gross", "gross amount"},
    "fee": {"fee", "fees", "processing fee"},
    "net": {"net", "net amount"},
    "transfer_id": {"transferid", "id"},
    "date_sent": {"datesent", "sent"},
    "amount_usd": {"amountusd", "usd"},
    "date_delivered": {"datedelivered", "delivered"},
    "rate": {"rate", "курс"},
    "reference": {"reference", "призначення", "memo"},
    "target_account": {"targetaccount", "target"},
    "funding_source": {"fundingsource", "джерело"},
}


class IngestError(Exception):
    pass


def _canon_header(h: str) -> str | None:
    key = re.sub(r"[\s_./-]+", "", str(h or "")).strip().lower()
    for canon, syns in _HEADER_SYNONYMS.items():
        if key == re.sub(r"[\s_./-]+", "", canon) or key in {
                re.sub(r"[\s_./-]+", "", s) for s in syns}:
            return canon
    return None


def _map_headers(row: list) -> dict[int, str]:
    mapping = {}
    for i, cell in enumerate(row):
        canon = _canon_header(cell)
        if canon:
            mapping[i] = canon
    return mapping


def _iso(value) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    s = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    raise IngestError(f"unparseable date: {value!r}")


def _num(value):
    if value is None or str(value).strip() == "":
        return None
    s = str(value).replace(" ", "").replace(" ", "")
    s = s.replace(",", ".") if s.count(",") == 1 and s.count(".") == 0 else s.replace(",", "")
    return money(s)


def _rows_from_sheet(ws) -> list[list]:
    return [[c for c in row] for row in ws.iter_rows(values_only=True)]


def _rows_from_csv(path: Path) -> list[list]:
    with open(path, encoding="utf-8-sig", newline="") as f:
        return [row for row in csv.reader(f)]


def _parse_table(rows: list[list], required: set[str], doc: Document,
                 locator_prefix: str) -> list[dict]:
    """Find the header row, then parse each data row into a dict of
    canonical fields + a SourceRef."""
    header_idx, mapping = None, {}
    for i, row in enumerate(rows[:10]):
        m = _map_headers(row)
        if required.issubset(set(m.values())):
            header_idx, mapping = i, m
            break
    if header_idx is None:
        raise IngestError(
            f"{doc.file}: could not find required columns {sorted(required)}")
    out = []
    for rownum, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if all(c is None or str(c).strip() == "" for c in row):
            continue
        rec = {mapping[i]: row[i] for i in mapping if i < len(row)}
        rec["_source"] = SourceRef(doc.doc_id, doc.file,
                                   f"{locator_prefix} row={rownum}")
        out.append(rec)
    return out


# ---------------------------------------------------------------------------
# Accounting workbook

_SHEET_HINTS = {
    "general_ledger": ("generalledger", "ledger", "журнал"),
    "payroll": ("payroll", "зарпл"),
    "chart": ("chartofaccounts", "chart", "план"),
    "balances": ("balances", "залишк"),
}


def read_workbook(path: Path, doc: Document) -> dict:
    """Return {'ledger': [...], 'payroll': [...], 'balances': [...],
    'chart': {...}, 'missing_sheets': [...]}"""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    found: dict[str, str] = {}
    for sheet in wb.sheetnames:
        key = re.sub(r"[\s_-]+", "", sheet).lower()
        for canon, hints in _SHEET_HINTS.items():
            if canon not in found and any(h in key for h in hints):
                found[canon] = sheet

    result = {"ledger": [], "payroll": [], "balances": [], "chart": {},
              "missing_sheets": []}
    for canon in ("general_ledger", "balances"):
        if canon not in found:
            result["missing_sheets"].append(canon)

    if "general_ledger" in found:
        rows = _rows_from_sheet(wb[found["general_ledger"]])
        for rec in _parse_table(rows, {"date", "amount_uah", "counterparty"},
                                doc, f"sheet={found['general_ledger']}"):
            rec["date"] = _iso(rec["date"])
            rec["amount_uah"] = _num(rec["amount_uah"])
            result["ledger"].append(rec)

    if "payroll" in found:
        rows = _rows_from_sheet(wb[found["payroll"]])
        for rec in _parse_table(rows, {"employee", "gross_uah"},
                                doc, f"sheet={found['payroll']}"):
            for k in ("gross_uah", "taxes_uah", "net_uah"):
                if k in rec:
                    rec[k] = _num(rec[k])
            result["payroll"].append(rec)

    if "balances" in found:
        rows = _rows_from_sheet(wb[found["balances"]])
        for rec in _parse_table(rows, {"bank_account", "opening_uah", "closing_uah"},
                                doc, f"sheet={found['balances']}"):
            rec["opening_uah"] = _num(rec["opening_uah"])
            rec["closing_uah"] = _num(rec["closing_uah"])
            result["balances"].append(rec)

    if "chart" in found:
        rows = _rows_from_sheet(wb[found["chart"]])
        for row in rows[1:]:
            if row and row[0]:
                result["chart"][str(row[0]).strip()] = str(row[1] or "").strip()
    wb.close()
    return result


# ---------------------------------------------------------------------------
# Bank statements

def read_bank_statement(path: Path, doc: Document) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        rows = _rows_from_csv(path)
    elif suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        rows = _rows_from_sheet(wb[wb.sheetnames[0]])
        wb.close()
    elif suffix == ".pdf":
        rows = _pdf_table_rows(path, doc)
        if rows is None:
            return []  # low-confidence, gap raised by caller via doc flag
    else:
        raise IngestError(f"{doc.file}: unsupported bank statement format {suffix}")

    out = []
    for rec in _parse_table(rows, {"date", "amount_uah"}, doc, "table"):
        rec["date"] = _iso(rec["date"])
        rec["amount_uah"] = _num(rec["amount_uah"])
        if "balance" in rec:
            rec["balance"] = _num(rec["balance"])
        out.append(rec)
    return out


# ---------------------------------------------------------------------------
# Center statistics

def read_center_stats(path: Path, doc: Document) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix in (".yaml", ".yml"):
        with open(path, encoding="utf-8") as f:
            data = yaml.safe_load(f) or {}
        out = []
        for i, rec in enumerate(data.get("centers", [])):
            rec["_source"] = SourceRef(doc.doc_id, doc.file, f"centers[{i}]")
            out.append(rec)
        return out
    if suffix == ".csv":
        rows = _rows_from_csv(path)
    elif suffix == ".pdf":
        rows = _pdf_table_rows(path, doc)
        if rows is None:
            return []
    else:
        raise IngestError(f"{doc.file}: unsupported stats format {suffix}")
    out = []
    for rec in _parse_table(rows, {"center", "children_served"}, doc, "table"):
        for k in ("children_served", "consultations", "classes", "new_children"):
            if k in rec and rec[k] is not None and str(rec[k]).strip() != "":
                rec[k] = int(float(str(rec[k])))
        out.append(rec)
    return out


# ---------------------------------------------------------------------------
# Donation platform exports (Zeffy / Stripe / PayPal) — USD

def read_platform_export(path: Path, doc: Document) -> list[dict]:
    if path.suffix.lower() != ".csv":
        raise IngestError(f"{doc.file}: platform exports must be CSV")
    out = []
    for rec in _parse_table(_rows_from_csv(path), {"date", "gross"},
                            doc, "table"):
        rec["date"] = _iso(rec["date"])
        for k in ("gross", "fee", "net"):
            if k in rec:
                rec[k] = _num(rec[k])
        if rec.get("net") is None and rec.get("gross") is not None:
            rec["net"] = rec["gross"] - (rec.get("fee") or 0)
        out.append(rec)
    return out


# ---------------------------------------------------------------------------
# Wise transfers

def read_wise(path: Path, doc: Document) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix in (".yaml", ".yml"):
        with open(path, encoding="utf-8") as f:
            data = yaml.safe_load(f) or {}
        out = []
        for i, rec in enumerate(data.get("transfers", [])):
            rec["_source"] = SourceRef(doc.doc_id, doc.file, f"transfers[{i}]")
            out.append(rec)
        return out
    if suffix != ".csv":
        raise IngestError(f"{doc.file}: unsupported wise format {suffix}")
    out = []
    for rec in _parse_table(_rows_from_csv(path),
                            {"transfer_id", "amount_usd"}, doc, "table"):
        rec["date_sent"] = _iso(rec["date_sent"]) if rec.get("date_sent") else None
        if rec.get("date_delivered"):
            rec["date_delivered"] = _iso(rec["date_delivered"])
        for k in ("amount_usd", "amount_uah", "rate"):
            if k in rec:
                rec[k] = _num(rec[k])
        out.append(rec)
    return out


# ---------------------------------------------------------------------------
# Funding requests (optional, PDF or text)

_AMOUNT_RE = re.compile(r"\$\s?([\d,]+(?:\.\d{2})?)")


def read_funding_request(path: Path, doc: Document) -> dict:
    text = ""
    if path.suffix.lower() == ".pdf":
        text = _pdf_text(path) or ""
    else:
        text = path.read_text(encoding="utf-8", errors="replace")
    amounts = [money(m.replace(",", "")) for m in _AMOUNT_RE.findall(text)]
    if not text.strip():
        doc.extraction_confidence = "low"
    return {
        "doc_id": doc.doc_id,
        "file": doc.file,
        "requested_usd": str(max(amounts)) if amounts else None,
        "text_excerpt": text[:400],
    }


# ---------------------------------------------------------------------------
# PDF helpers — native text only; scanned PDFs become low-confidence gaps

def _pdf_text(path: Path) -> str | None:
    import pdfplumber
    with pdfplumber.open(path) as pdf:
        text = "\n".join((page.extract_text() or "") for page in pdf.pages)
    return text if text.strip() else None


def _pdf_table_rows(path: Path, doc: Document) -> list[list] | None:
    import pdfplumber
    rows: list[list] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables():
                rows.extend(table)
    if not rows:
        doc.extraction_confidence = "low"
        doc.notes = "no extractable table/text — scanned document needs OCR or a CSV re-upload"
        return None
    return rows
